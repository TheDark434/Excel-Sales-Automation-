from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference


file_path = Path(
    r"D:\Projects\Excel Sales Automation\output\sales_report.xlsx"
)


def create_branch_revenue_chart():

    workbook = load_workbook(file_path)

    branch_sheet = workbook["Branch Analysis"]

    # Create branch revenue chart
    chart = BarChart()

    chart.title = "Revenue by Branch"
    chart.y_axis.title = "Revenue"
    chart.x_axis.title = "Branch"

    data = Reference(
        branch_sheet,
        min_col=2,
        min_row=1,
        max_row=branch_sheet.max_row
    )

    categories = Reference(
        branch_sheet,
        min_col=1,
        min_row=2,
        max_row=branch_sheet.max_row
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    branch_sheet.add_chart(chart, "D2")

    workbook.save(file_path)

    print("Branch revenue chart created successfully!")


def create_product_quantity_chart():

    workbook = load_workbook(file_path)

    product_sheet = workbook["Product Analysis"]

    # Create product quantity chart
    chart = BarChart()

    chart.title = "Quantity Sold by Product"
    chart.y_axis.title = "Quantity Sold"
    chart.x_axis.title = "Product"

    data = Reference(
        product_sheet,
        min_col=2,
        min_row=1,
        max_row=product_sheet.max_row
    )

    categories = Reference(
        product_sheet,
        min_col=1,
        min_row=2,
        max_row=product_sheet.max_row
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    product_sheet.add_chart(chart, "E2")

    workbook.save(file_path)

    print("Product quantity chart created successfully!")

def create_category_revenue_chart():

    workbook = load_workbook(file_path)

    category_sheet = workbook["Category Analysis"]

    # Create category revenue chart
    chart = BarChart()

    chart.title = "Revenue by Category"
    chart.y_axis.title = "Revenue"
    chart.x_axis.title = "Category"

    data = Reference(
        category_sheet,
        min_col=2,
        min_row=1,
        max_row=category_sheet.max_row
    )

    categories = Reference(
        category_sheet,
        min_col=1,
        min_row=2,
        max_row=category_sheet.max_row
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    category_sheet.add_chart(chart, "D2")

    workbook.save(file_path)

    print("Category revenue chart created successfully!")


 

   
