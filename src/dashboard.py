from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter


def create_dashboard():

    file_name = r"D:\Projects\Excel Sales Automation\output\sales_report.xlsx"

    workbook = load_workbook(file_name)
    summary_sheet = workbook["Summary"]

    # Clear the existing Summary sheet
    for row in summary_sheet.iter_rows():
        for cell in row:
            cell.value = None
            

    # Dashboard title
    summary_sheet.merge_cells("A1:H1")
    summary_sheet["A1"] = "SALES DASHBOARD"
    summary_sheet["A1"].font = Font(
        bold=True,
        size=18,
        color="FFFFFF"
    )
    summary_sheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )
    summary_sheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )
    summary_sheet.row_dimensions[1].height = 30

    # KPI labels
    summary_sheet["A3"] = "Total Revenue"
    summary_sheet["D3"] = "Total Quantity"
    summary_sheet["G3"] = "Average Transaction Value"

    # KPI values
    summary_sheet["A4"] = 1679920
    summary_sheet["D4"] = 192
    summary_sheet["G4"] = 10059.40

    # KPI styling
    kpi_labels = ["A3", "D3", "G3"]
    kpi_values = ["A4", "D4", "G4"]

    for cell in kpi_labels:
        summary_sheet[cell].font = Font(
            bold=True,
            color="FFFFFF"
        )
        summary_sheet[cell].fill = PatternFill(
            fill_type="solid",
            fgColor="5B9BD5"
        )
        summary_sheet[cell].alignment = Alignment(
            horizontal="center"
        )

    for cell in kpi_values:
        summary_sheet[cell].font = Font(
            bold=True,
            size=14
        )
        summary_sheet[cell].alignment = Alignment(
            horizontal="center"
        )

    # Number formats
    summary_sheet["A4"].number_format = '₹#,##0.00'
    summary_sheet["D4"].number_format = '#,##0'
    summary_sheet["G4"].number_format = '₹#,##0.00'

    # Borders
    thin_side = Side(style="thin")

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    for cell in kpi_labels + kpi_values:
        summary_sheet[cell].border = border

    # Dashboard section titles
    summary_sheet["A6"] = "Revenue by Branch"
    summary_sheet["E6"] = "Quantity by Product"

    for cell in ["A6", "E6"]:
        summary_sheet[cell].font = Font(
            bold=True,
            size=12
        )

    # Branch chart
    branch_sheet = workbook["Branch Analysis"]

    branch_chart = BarChart()
    branch_chart.title = "Revenue by Branch"
    branch_chart.y_axis.title = "Revenue"
    branch_chart.x_axis.title = "Branch"

    branch_data = Reference(
        branch_sheet,
        min_col=2,
        min_row=1,
        max_row=branch_sheet.max_row
    )

    branch_categories = Reference(
        branch_sheet,
        min_col=1,
        min_row=2,
        max_row=branch_sheet.max_row
    )

    branch_chart.add_data(
        branch_data,
        titles_from_data=True
    )
    branch_chart.set_categories(branch_categories)

    branch_chart.height = 8
    branch_chart.width = 12

    summary_sheet.add_chart(branch_chart, "A7")

    # Product chart
    product_sheet = workbook["Product Analysis"]

    product_chart = BarChart()
    product_chart.title = "Quantity Sold by Product"
    product_chart.y_axis.title = "Quantity"
    product_chart.x_axis.title = "Product"

    product_data = Reference(
        product_sheet,
        min_col=2,
        min_row=1,
        max_row=product_sheet.max_row
    )

    product_categories = Reference(
        product_sheet,
        min_col=1,
        min_row=2,
        max_row=product_sheet.max_row
    )

    product_chart.add_data(
        product_data,
        titles_from_data=True
    )
    product_chart.set_categories(product_categories)

    product_chart.height = 8
    product_chart.width = 12

    summary_sheet.add_chart(product_chart, "E7")

    # Column widths
    for column in range(1, 9):
        summary_sheet.column_dimensions[
            get_column_letter(column)
        ].width = 18

    # Freeze panes
    summary_sheet.freeze_panes = "A3"

    # Save workbook
    workbook.save(file_name)

    print("Dashboard created successfully!")



